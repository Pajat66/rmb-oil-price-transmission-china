# -*- coding: utf-8 -*-
"""
生成长样本主表：1998-01 ~ 2026-02
依赖：pandas, openpyxl
注意：CPI 四个文件需为 .xlsx（因为 openpyxl 不能读 .xls）
"""
import pandas as pd
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# === 1) Resolve the repository root from this script's location ===
PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW = PROJECT_ROOT / "01_raw"
OUT = PROJECT_ROOT / "02_clean" / "monthly_master" / "master_monthly_long_1998_2026.xlsx"

FILES = {
    "eia": RAW / "01_eia_brent" / "EIA_brent_usd_monthly_1987-latest_2026.3.22.csv",
    "fred": RAW / "02_fx_pbc_fred" / "FRED_exchus_cny_per_usd_monthly_1981-latest_2026.3.22.xlsx",
    "cpi_98_15": RAW / "03_nbs_price" / "98年-15年月度消费价格分类指数.xlsx",
    "cpi_16_20": RAW / "03_nbs_price" / "16年-20年月度消费价格分类指数.xlsx",
    "cpi_21_25": RAW / "03_nbs_price" / "21年-25年月度消费价格分类指数.xlsx",
    "cpi_26": RAW / "03_nbs_price" / "26年1-2月月度消费价格分类指数.xlsx",
    "ppi": RAW / "03_nbs_price" / "98年-26年2月工业生产者出厂价格指数.csv",
    "ppi_purchase": RAW / "03_nbs_price" / "98年-26年2月工业生产者购入价格指数.csv",
    "iva": RAW / "03_nbs_price" / "98年-26年2月规上工业增加值增长速度.csv",
}

def ym_to_str(s):
    if pd.isna(s): return None
    s = str(s).strip()
    m = re.match(r"(\d{4})年(\d{1,2})月", s)
    if m: return f"{m.group(1)}-{int(m.group(2)):02d}"
    try: return pd.to_datetime(s).strftime("%Y-%m")
    except: return None

def read_wide_excel_xlsx(path: Path):
    df = pd.read_excel(path, header=None, engine="openpyxl")
    headers = df.iloc[2].tolist()
    df2 = df.iloc[3:].copy()
    df2.columns = headers
    return df2[df2["指标"].notna()].copy()

def read_wide_csv(path: Path):
    df = pd.read_csv(path, encoding="gbk", skiprows=2)
    return df[df["指标"].notna()].copy()

def extract_row(df, labels):
    if isinstance(labels, str): labels = [labels]
    for label in labels:
        row = df[df["指标"] == label]
        if not row.empty:
            row = row.iloc[0]
            out = []
            for col in df.columns[1:]:
                month = ym_to_str(col)
                if month is None: continue
                val = row[col]
                if pd.notna(val): out.append((month, val))
            return pd.DataFrame(out, columns=["month","value"]).sort_values("month")
    return pd.DataFrame(columns=["month","value"])

def main():
    months = pd.period_range("1998-01", "2026-02", freq="M").strftime("%Y-%m")
    master = pd.DataFrame({"month": months})

    # Brent
    eia = pd.read_csv(FILES["eia"], skiprows=4)
    eia["month"] = pd.to_datetime(eia["Month"], format="%b %Y").dt.strftime("%Y-%m")
    price_col = [c for c in eia.columns if c != "Month"][0]
    eia = eia[["month", price_col]].rename(columns={price_col:"brent_usd"})
    master = master.merge(eia, on="month", how="left")

    # FX
    fred = pd.read_excel(FILES["fred"], sheet_name="Monthly", engine="openpyxl")
    fred["month"] = pd.to_datetime(fred["observation_date"]).dt.strftime("%Y-%m")
    fred = fred[["month","EXCHUS"]].rename(columns={"EXCHUS":"cny_per_usd"})
    master = master.merge(fred, on="month", how="left")

    # CPI (4 periods)
    cpi_dfs = {
        "98_15": read_wide_excel_xlsx(FILES["cpi_98_15"]),
        "16_20": read_wide_excel_xlsx(FILES["cpi_16_20"]),
        "21_25": read_wide_excel_xlsx(FILES["cpi_21_25"]),
        "26": read_wide_excel_xlsx(FILES["cpi_26"]),
    }
    labelsets = {
        "cpi_yoy": [["居民消费价格指数(上年同月=100)"]],
        "cpi_transport_comm_yoy": [
            ["交通和通讯类居民消费价格指数(上年同月=100)"],
            ["交通和通信类居民消费价格指数(上年同月=100)"],
            ["交通通信类居民消费价格指数(上年同月=100)"],
        ],
        "core_cpi_proxy_yoy": [["不包括食品和能源居民消费价格指数(上年同月=100)"]],
    }
    for out_col, alias_groups in labelsets.items():
        parts = []
        for df in cpi_dfs.values():
            found = pd.DataFrame(columns=["month","value"])
            for labels in alias_groups:
                found = extract_row(df, labels)
                if not found.empty: break
            if not found.empty: parts.append(found)
        if parts:
            series = pd.concat(parts, ignore_index=True).drop_duplicates("month", keep="first").sort_values("month")
            series = series.rename(columns={"value": out_col})
            master = master.merge(series, on="month", how="left")
        else:
            master[out_col] = pd.NA

    # PPI
    ppi_df = read_wide_csv(FILES["ppi"])
    ppip_df = read_wide_csv(FILES["ppi_purchase"])
    ppi = extract_row(ppi_df, "工业生产者出厂价格指数(上年同月=100)").rename(columns={"value":"ppi_yoy"})
    ppi_purchase = extract_row(ppip_df, "工业生产者购进价格指数(上年同月=100)").rename(columns={"value":"ppi_purchase_yoy"})
    ppi_fuel = extract_row(ppip_df, "燃料、动力类购进价格指数(上年同月=100)").rename(columns={"value":"ppi_fuel_power_yoy"})
    master = master.merge(ppi, on="month", how="left").merge(ppi_purchase, on="month", how="left").merge(ppi_fuel, on="month", how="left")

    # Industrial VA
    iva_df = read_wide_csv(FILES["iva"])
    iva_yoy = extract_row(iva_df, "规上工业增加值同比增长(%)").rename(columns={"value":"indust_va_yoy"})
    iva_cum = extract_row(iva_df, "规上工业增加值累计增长(%)").rename(columns={"value":"indust_va_cum_yoy"})
    master = master.merge(iva_yoy, on="month", how="left").merge(iva_cum, on="month", how="left")

    # === write workbook ===
    headers = [
        "month","brent_usd","cny_per_usd","oil_rmb",
        "cpi_yoy","cpi_transport_comm_yoy","core_cpi_proxy_yoy",
        "ppi_yoy","ppi_purchase_yoy","ppi_fuel_power_yoy",
        "indust_va_yoy","indust_va_cum_yoy"
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "master_monthly_long"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for c,h in enumerate(headers, start=1):
        cell = ws.cell(1,c,value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for r, row in enumerate(master.itertuples(index=False), start=2):
        ws.cell(r,1,value=row.month)
        ws.cell(r,2,value=None if pd.isna(row.brent_usd) else float(row.brent_usd))
        ws.cell(r,3,value=None if pd.isna(row.cny_per_usd) else float(row.cny_per_usd))
        ws.cell(r,4,value=f'=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})')
        vals = [
            row.cpi_yoy,row.cpi_transport_comm_yoy,row.core_cpi_proxy_yoy,
            row.ppi_yoy,row.ppi_purchase_yoy,row.ppi_fuel_power_yoy,
            row.indust_va_yoy,row.indust_va_cum_yoy
        ]
        for i,v in enumerate(vals, start=5):
            ws.cell(r,i,value=None if pd.isna(v) else float(v))

    # formats
    for r in range(2, ws.max_row+1):
        for c in [2,3,4]:
            ws.cell(r,c).number_format = "0.0000"
        for c in range(5,13):
            ws.cell(r,c).number_format = "0.0"

    widths = {1:10,2:12,3:12,4:12,5:12,6:20,7:20,8:12,9:18,10:20,11:14,12:16}
    for col,w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("已生成：", OUT)

if __name__ == "__main__":
    main()
