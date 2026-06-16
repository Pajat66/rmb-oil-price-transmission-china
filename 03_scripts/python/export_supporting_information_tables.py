from pathlib import Path
from math import erfc, sqrt

import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "04_results" / "supporting_information"
OUT_FILE = OUT_DIR / "S1_Table.xlsx"
MAPPING_FILE = OUT_DIR / "supporting_table_mapping.csv"
CAPTIONS_FILE = OUT_DIR / "suggested_supporting_information_captions.txt"

LONG_FILE = ROOT / "04_results" / "chapter2" / "master_long_clean_v1.xlsx"
SHORT_FILE = ROOT / "04_results" / "chapter2" / "master_short_clean_v1.xlsx"
AUX_LONG_FILE = ROOT / "02_clean" / "monthly_master" / "master_monthly_long_1998_2026.xlsx"
LONG_MAIN_VARS = [
    "brent_usd",
    "cny_per_usd",
    "oil_rmb",
    "cpi_yoy",
    "cpi_transport_comm_yoy",
    "core_cpi_proxy_yoy",
    "ppi_yoy",
    "ppi_purchase_yoy",
    "ppi_fuel_power_yoy",
    "indust_va_yoy",
    "indust_va_cum_yoy",
]

UNIT_ROOT_VARS = [
    "brent_usd",
    "cny_per_usd",
    "oil_rmb",
    "cpi_yoy",
    "cpi_transport_comm_yoy",
    "ppi_yoy",
    "ppi_purchase_yoy",
    "ppi_fuel_power_yoy",
    "indust_va_yoy",
    "indust_va_cum_yoy",
]

SHORT_MAIN_VARS = [
    "brent_usd",
    "cny_per_usd",
    "oil_rmb",
    "cpi_yoy",
    "cpi_transport_comm_yoy",
    "core_cpi_proxy_yoy",
    "ppi_yoy",
    "ppi_purchase_yoy",
    "ppi_fuel_power_yoy",
    "import_crude_oil_qty_10k_ton",
    "import_crude_oil_value_100m_rmb",
    "import_crude_oil_unit_price_yuan_per_kg",
]

VARIABLE_METADATA = {
    "month": {
        "abbreviation": "Month",
        "description": "Calendar month",
        "unit": "YYYY-MM",
        "source_or_derivation": "Unified monthly date index",
    },
    "brent_usd": {
        "abbreviation": "Brent",
        "description": "Brent crude oil price",
        "unit": "USD per barrel",
        "source_or_derivation": "EIA Europe Brent spot price",
    },
    "cny_per_usd": {
        "abbreviation": "CNY/USD",
        "description": "RMB per US dollar exchange rate",
        "unit": "CNY per USD",
        "source_or_derivation": "FRED monthly exchange rate series",
    },
    "oil_rmb": {
        "abbreviation": "RMB oil price",
        "description": "Brent price converted into RMB",
        "unit": "CNY per barrel",
        "source_or_derivation": "brent_usd multiplied by cny_per_usd",
    },
    "cpi_yoy": {
        "abbreviation": "CPI YoY",
        "description": "Consumer price index, year on year",
        "unit": "Index, previous year same month equals 100",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "cpi_transport_comm_yoy": {
        "abbreviation": "Transport CPI YoY",
        "description": "Transport and communication CPI, year on year",
        "unit": "Index, previous year same month equals 100",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "core_cpi_proxy_yoy": {
        "abbreviation": "Core CPI proxy YoY",
        "description": "Proxy for core CPI, year on year",
        "unit": "Index, previous year same month equals 100",
        "source_or_derivation": "Constructed from available CPI components",
    },
    "ppi_yoy": {
        "abbreviation": "PPI YoY",
        "description": "Producer price index, year on year",
        "unit": "Index, previous year same month equals 100",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "ppi_purchase_yoy": {
        "abbreviation": "Purchase PPI YoY",
        "description": "Purchasing price index for industrial producers, year on year",
        "unit": "Index, previous year same month equals 100",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "ppi_fuel_power_yoy": {
        "abbreviation": "Fuel and power PPI YoY",
        "description": "Purchase price index for fuel and power, year on year",
        "unit": "Index, previous year same month equals 100",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "indust_va_yoy": {
        "abbreviation": "Industrial VA YoY",
        "description": "Industrial value added, year on year",
        "unit": "Percent",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "indust_va_cum_yoy": {
        "abbreviation": "Cumulative industrial VA YoY",
        "description": "Cumulative industrial value added, year on year",
        "unit": "Percent",
        "source_or_derivation": "National Bureau of Statistics of China",
    },
    "import_crude_oil_qty_10k_ton": {
        "abbreviation": "Crude import volume",
        "description": "Crude oil import volume",
        "unit": "Ten-thousand tons",
        "source_or_derivation": "China Customs monthly import data",
    },
    "import_crude_oil_value_100m_rmb": {
        "abbreviation": "Crude import value",
        "description": "Crude oil import value",
        "unit": "Hundred-million CNY",
        "source_or_derivation": "China Customs monthly import data",
    },
    "import_crude_oil_unit_price_yuan_per_kg": {
        "abbreviation": "Crude import unit price",
        "description": "Crude oil import unit price",
        "unit": "CNY per kg",
        "source_or_derivation": "Import value divided by import volume",
    },
}


def stars(p_value):
    if p_value is None or not np.isfinite(p_value):
        return ""
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def normal_two_sided_pvalue(z_value):
    if z_value is None or not np.isfinite(z_value):
        return np.nan
    return erfc(abs(float(z_value)) / sqrt(2.0))


def ols_hc3(y, X):
    y = np.asarray(y, dtype=float)
    X = np.asarray(X, dtype=float)
    xtx_inv = np.linalg.pinv(X.T @ X)
    beta = xtx_inv @ X.T @ y
    fitted = X @ beta
    resid = y - fitted
    nobs, nparams = X.shape
    hat = np.sum((X @ xtx_inv) * X, axis=1)
    safe_denom = np.maximum(1.0 - hat, 1e-12)
    scaled = (resid / safe_denom) ** 2
    meat = X.T @ (scaled[:, None] * X)
    cov_hc3 = xtx_inv @ meat @ xtx_inv
    se_hc3 = np.sqrt(np.maximum(np.diag(cov_hc3), 0))

    ss_res = float(np.sum(resid ** 2))
    ss_tot = float(np.sum((y - y.mean()) ** 2))
    r2 = 1.0 - ss_res / ss_tot if ss_tot > 0 else np.nan
    adj_r2 = 1.0 - (1.0 - r2) * (nobs - 1) / (nobs - nparams) if nobs > nparams and np.isfinite(r2) else np.nan
    return {
        "beta": beta,
        "cov": cov_hc3,
        "se": se_hc3,
        "pvalues": np.array([normal_two_sided_pvalue(b / s) if s > 0 else np.nan for b, s in zip(beta, se_hc3)]),
        "adj_r2": adj_r2,
        "nobs": nobs,
    }


def ols_standard(y, X):
    y = np.asarray(y, dtype=float)
    X = np.asarray(X, dtype=float)
    xtx_inv = np.linalg.pinv(X.T @ X)
    beta = xtx_inv @ X.T @ y
    resid = y - X @ beta
    nobs, nparams = X.shape
    sigma2 = float(np.sum(resid ** 2) / max(nobs - nparams, 1))
    cov = sigma2 * xtx_inv
    se = np.sqrt(np.maximum(np.diag(cov), 0))
    return beta, se, resid


def approximate_adf_pvalue(stat):
    if not np.isfinite(stat):
        return np.nan
    points = [(-3.45, 0.01), (-2.87, 0.05), (-2.57, 0.10)]
    if stat <= points[0][0]:
        return 0.01
    if stat >= points[-1][0]:
        return min(0.99, 0.10 + (stat - points[-1][0]) * 0.15)
    for (x0, p0), (x1, p1) in zip(points[:-1], points[1:]):
        if x0 <= stat <= x1:
            weight = (stat - x0) / (x1 - x0)
            return p0 + weight * (p1 - p0)
    return np.nan


def adf_constant_aic(series):
    values = pd.Series(series).dropna().astype(float).to_numpy()
    n = values.shape[0]
    max_lag = min(12, max(0, n // 4))
    best = None
    for lag in range(max_lag + 1):
        dy = np.diff(values)
        rows = []
        target = []
        for t in range(lag + 1, n):
            row = [1.0, values[t - 1]]
            for j in range(1, lag + 1):
                row.append(dy[t - j - 1])
            rows.append(row)
            target.append(dy[t - 1])
        if len(target) <= len(rows[0]):
            continue
        X = np.asarray(rows, dtype=float)
        y = np.asarray(target, dtype=float)
        beta, se, resid = ols_standard(y, X)
        rss = float(np.sum(resid ** 2))
        nobs = y.shape[0]
        k = X.shape[1]
        aic = np.log(max(rss / nobs, 1e-12)) + 2 * k / nobs
        t_stat = beta[1] / se[1] if se[1] > 0 else np.nan
        if best is None or aic < best["aic"]:
            best = {
                "stat": float(t_stat),
                "pvalue": float(approximate_adf_pvalue(t_stat)),
                "used_lag": lag,
                "nobs": int(nobs),
                "aic": float(aic),
            }
    return best


def load_master(path):
    df = pd.read_excel(path, engine="openpyxl")
    df["month"] = pd.to_datetime(df["month"], errors="coerce")
    df = df.sort_values("month").reset_index(drop=True)
    for col in df.columns:
        if col != "month":
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if {"brent_usd", "cny_per_usd"}.issubset(df.columns):
        computed_oil_rmb = df["brent_usd"] * df["cny_per_usd"]
        if "oil_rmb" not in df.columns:
            df["oil_rmb"] = computed_oil_rmb
        else:
            df["oil_rmb"] = df["oil_rmb"].where(df["oil_rmb"].notna(), computed_oil_rmb)
    return df


def add_industrial_control(df):
    if "indust_va_cum_yoy" in df.columns:
        return df
    aux = load_master(AUX_LONG_FILE)[["month", "indust_va_yoy", "indust_va_cum_yoy"]]
    out = df.merge(aux, on="month", how="left")
    return out


def variable_table():
    rows = []
    ordered_variables = ["month"] + sorted(set(LONG_MAIN_VARS + SHORT_MAIN_VARS), key=(LONG_MAIN_VARS + SHORT_MAIN_VARS).index)
    for var in ordered_variables:
        metadata = VARIABLE_METADATA.get(
            var,
            {
                "abbreviation": var,
                "description": "",
                "unit": "",
                "source_or_derivation": "",
            },
        )
        sample_scope = []
        if var in LONG_MAIN_VARS or var == "month":
            sample_scope.append("Long sample")
        if var in SHORT_MAIN_VARS or var == "month":
            sample_scope.append("Short sample")
        rows.append(
            {
                "Variable name": var,
                "Abbreviation": metadata["abbreviation"],
                "Description": metadata["description"],
                "Unit": metadata["unit"],
                "Source or derivation": metadata["source_or_derivation"],
                "Sample scope": "; ".join(sample_scope),
            }
        )
    return pd.DataFrame(rows)


def descriptive_stats(df, variables):
    rows = []
    for var in variables:
        if var not in df.columns:
            continue
        x = df[var].dropna()
        observed_months = df.loc[df[var].notna(), "month"]
        metadata = VARIABLE_METADATA.get(var, {"abbreviation": var, "description": ""})
        rows.append(
            {
                "Variable name": var,
                "Abbreviation": metadata["abbreviation"],
                "Description": metadata["description"],
                "N": int(x.shape[0]),
                "Mean": round(float(x.mean()), 4) if not x.empty else np.nan,
                "Std. dev.": round(float(x.std(ddof=1)), 4) if x.shape[0] > 1 else np.nan,
                "Min": round(float(x.min()), 4) if not x.empty else np.nan,
                "P25": round(float(x.quantile(0.25)), 4) if not x.empty else np.nan,
                "Median": round(float(x.median()), 4) if not x.empty else np.nan,
                "P75": round(float(x.quantile(0.75)), 4) if not x.empty else np.nan,
                "Max": round(float(x.max()), 4) if not x.empty else np.nan,
                "First observed month": observed_months.min().strftime("%Y-%m") if not observed_months.empty else "",
                "Last observed month": observed_months.max().strftime("%Y-%m") if not observed_months.empty else "",
            }
        )
    return pd.DataFrame(rows)


def adf_table(df, variables, difference=False):
    rows = []
    for var in variables:
        if var not in df.columns:
            continue
        series = df[var].dropna()
        if difference:
            series = series.diff().dropna()
        metadata = VARIABLE_METADATA.get(var, {"abbreviation": var, "description": ""})
        if series.shape[0] < 12 or np.isclose(series.std(ddof=1), 0):
            rows.append(
                {
                    "Variable name": var,
                    "Abbreviation": metadata["abbreviation"],
                    "Description": metadata["description"],
                    "Transformation": "First difference" if difference else "Level",
                    "N": int(series.shape[0]),
                    "ADF statistic": np.nan,
                    "Approximate p-value": np.nan,
                    "Used lags": np.nan,
                    "Critical value 1%": np.nan,
                    "Critical value 5%": np.nan,
                    "Critical value 10%": np.nan,
                    "Stationary at 5%": "",
                    "Note": "Insufficient non-missing observations or no variation",
                }
            )
            continue
        result = adf_constant_aic(series)
        stat = result["stat"]
        p_value = result["pvalue"]
        used_lag = result["used_lag"]
        nobs = result["nobs"]
        crit = {"1%": -3.45, "5%": -2.87, "10%": -2.57}
        rows.append(
            {
                "Variable name": var,
                "Abbreviation": metadata["abbreviation"],
                "Description": metadata["description"],
                "Transformation": "First difference" if difference else "Level",
                "N": int(nobs),
                "ADF statistic": round(float(stat), 4),
                "Approximate p-value": round(float(p_value), 4),
                "Used lags": int(used_lag),
                "Critical value 1%": round(float(crit["1%"]), 4),
                "Critical value 5%": round(float(crit["5%"]), 4),
                "Critical value 10%": round(float(crit["10%"]), 4),
                "Stationary at 5%": "Yes" if p_value < 0.05 else "No",
                "Note": "ADF with constant; lag length selected by AIC; p-value is approximated from conventional critical values",
            }
        )
    return pd.DataFrame(rows)


def fit_lagged_model(df, yvar, xvars, lags=6, controls=None, min_obs=20):
    controls = controls or []
    temp = df[["month", yvar] + xvars + controls].copy()
    predictor_cols = []
    for xvar in xvars:
        for lag in range(lags + 1):
            col = f"{xvar}_lag{lag}"
            temp[col] = temp[xvar].shift(lag)
            predictor_cols.append(col)
    predictor_cols.extend(controls)
    model_df = temp[["month", yvar] + predictor_cols].dropna()
    if model_df.shape[0] < min_obs:
        return None
    X = np.column_stack([np.ones(model_df.shape[0]), model_df[predictor_cols].to_numpy(dtype=float)])
    y = model_df[yvar]
    return {
        "fit": ols_hc3(y, X),
        "predictor_cols": predictor_cols,
        "model_df": model_df,
    }


def cumulative_row(model_result, xvar, max_lag, label):
    predictor_cols = model_result["predictor_cols"]
    model = model_result["fit"]
    idx = [1 + predictor_cols.index(f"{xvar}_lag{lag}") for lag in range(max_lag + 1)]
    effect = float(model["beta"][idx].sum())
    cov = model["cov"][np.ix_(idx, idx)]
    se = float(np.sqrt(max(cov.sum(), 0)))
    if se > 0:
        lower = effect - 1.96 * se
        upper = effect + 1.96 * se
        p_value = normal_two_sided_pvalue(effect / se)
    else:
        lower = upper = p_value = np.nan
    return {
        "Term": label,
        "Coefficient": round(effect, 4),
        "Std. error": round(se, 4) if np.isfinite(se) else np.nan,
        "95% CI lower": round(lower, 4) if np.isfinite(lower) else np.nan,
        "95% CI upper": round(upper, 4) if np.isfinite(upper) else np.nan,
        "p-value": round(p_value, 4) if np.isfinite(p_value) else np.nan,
        "Significance": stars(p_value),
        "Adjusted R2": round(float(model["adj_r2"]), 4),
        "N": int(model["nobs"]),
    }


def cumulative_effect_table(df):
    df = add_industrial_control(df)
    control_var = "indust_va_cum_yoy"
    targets = [
        ("ppi_fuel_power_yoy", "Purchase price index for fuel and power"),
        ("ppi_yoy", "Producer price index"),
        ("cpi_yoy", "Consumer price index"),
    ]
    rows = []
    for yvar, y_label in targets:
        model_result = fit_lagged_model(df, yvar, ["oil_rmb"], lags=6, controls=[control_var], min_obs=100)
        if model_result is None:
            continue
        for horizon in range(7):
            row = cumulative_row(model_result, "oil_rmb", horizon, f"Cumulative oil_rmb effect, lags 0-{horizon}")
            rows.append(
                {
                    "Dependent variable": yvar,
                    "Dependent variable label": y_label,
                    "Shock variable": "oil_rmb",
                    "Model": "6-period DLM with industrial value-added control",
                    "Control variable": control_var,
                    "Lag window": f"0-{horizon}",
                    "Cumulative effect": row["Coefficient"],
                    "Std. error": row["Std. error"],
                    "95% CI lower": row["95% CI lower"],
                    "95% CI upper": row["95% CI upper"],
                    "p-value": row["p-value"],
                    "Significance": row["Significance"],
                    "Adjusted R2": row["Adjusted R2"],
                    "N": row["N"],
                }
            )
    return pd.DataFrame(rows)


def oil_to_import_price_table(df):
    yvar = "import_crude_oil_unit_price_yuan_per_kg"
    rows = []
    model_specs = [
        ("Baseline DLM", []),
        ("DLM controlling for crude oil import quantity", ["import_crude_oil_qty_10k_ton"]),
    ]
    for model_name, controls in model_specs:
        model_result = fit_lagged_model(df, yvar, ["oil_rmb"], lags=3, controls=controls, min_obs=20)
        if model_result is None:
            continue
        model = model_result["fit"]
        predictor_cols = model_result["predictor_cols"]
        control_label = ", ".join(controls) if controls else "No additional control"
        for lag in range(4):
            idx = 1 + predictor_cols.index(f"oil_rmb_lag{lag}")
            coef = float(model["beta"][idx])
            se = float(model["se"][idx])
            p_value = float(model["pvalues"][idx])
            rows.append(
                {
                    "Model": model_name,
                    "Dependent variable": yvar,
                    "Control variable": control_label,
                    "Term": f"oil_rmb lag {lag}",
                    "Coefficient": round(coef, 4),
                    "Std. error": round(se, 4),
                    "95% CI lower": round(coef - 1.96 * se, 4),
                    "95% CI upper": round(coef + 1.96 * se, 4),
                    "p-value": round(p_value, 4),
                    "Significance": stars(p_value),
                    "Adjusted R2": round(float(model["adj_r2"]), 4),
                    "N": int(model["nobs"]),
                }
            )
        cumulative = cumulative_row(model_result, "oil_rmb", 3, "Cumulative oil_rmb effect, lags 0-3")
        cumulative.update(
            {
                "Model": model_name,
                "Dependent variable": yvar,
                "Control variable": control_label,
                "Term": "Cumulative oil_rmb effect, lags 0-3",
            }
        )
        rows.append(cumulative)
    return pd.DataFrame(rows)


def supporting_table_mapping():
    rows = [
        {
            "Sheet": "S1 Table",
            "Caption": "Variable names and abbreviations.",
            "Primary source": "03_scripts/python/export_supporting_information_tables.py",
            "Input data or model": "English variable metadata in the SI export script",
            "Notes": "Pure-English variable dictionary used by all SI sheets.",
        },
        {
            "Sheet": "S2 Table",
            "Caption": "Descriptive statistics of main variables for the long sample.",
            "Primary source": "04_results/chapter2/master_long_clean_v1.xlsx",
            "Input data or model": "Long-sample master data; oil_rmb is filled as brent_usd multiplied by cny_per_usd when the stored column is blank.",
            "Notes": "Includes core descriptive variables; core_cpi_proxy_yoy remains a supplementary descriptive variable.",
        },
        {
            "Sheet": "S3 Table",
            "Caption": "Descriptive statistics of main variables for the short sample.",
            "Primary source": "04_results/chapter2/master_short_clean_v1.xlsx",
            "Input data or model": "Short-sample master data.",
            "Notes": "Includes import cost variables used in mechanism tables.",
        },
        {
            "Sheet": "S4 Table",
            "Caption": "Unit root test results for level series.",
            "Primary source": "04_results/chapter2/master_long_clean_v1.xlsx",
            "Input data or model": "ADF tests with constant and AIC lag selection on core level series.",
            "Notes": "core_cpi_proxy_yoy is excluded from the core unit-root sheet because it is supplementary and not a core ARDL variable.",
        },
        {
            "Sheet": "S5 Table",
            "Caption": "Unit root test results for first-differenced series.",
            "Primary source": "04_results/chapter2/master_long_clean_v1.xlsx",
            "Input data or model": "ADF tests with constant and AIC lag selection on first differences of core series.",
            "Notes": "core_cpi_proxy_yoy is excluded from the core unit-root sheet because it is supplementary and not a core ARDL variable.",
        },
        {
            "Sheet": "S6 Table",
            "Caption": "Cumulative lag effects of the RMB-denominated oil price.",
            "Primary source": "04_results/chapter2/master_short_clean_v1.xlsx; 02_clean/monthly_master/master_monthly_long_1998_2026.xlsx",
            "Input data or model": "6-period DLM with lags 0-6 of oil_rmb and indust_va_cum_yoy control; cumulative effects are computed over lag windows.",
            "Notes": "The 0-3 rows align with the manuscript Table 4 cumulative effects and use N=117.",
        },
        {
            "Sheet": "S7 Table",
            "Caption": "Impacts of the RMB-denominated oil price on the crude oil import unit price.",
            "Primary source": "04_results/chapter2/master_short_clean_v1.xlsx",
            "Input data or model": "Baseline DLM with lags 0-3 of oil_rmb, plus the same DLM controlling for crude oil import quantity.",
            "Notes": "Reports lag terms and the 0-3 cumulative oil-price effect on import unit price, with and without import quantity control.",
        },
        {
            "Sheet": "S8 Table",
            "Caption": "Impacts of import costs and upstream/midstream prices on CPI YoY.",
            "Primary source": "04_results/chapter2/master_short_clean_v1.xlsx; 02_clean/monthly_master/master_monthly_long_1998_2026.xlsx",
            "Input data or model": "6-period CPI DLMs for import cost, fuel and power purchase price, and joint PPI plus fuel and power terms, all with indust_va_cum_yoy control.",
            "Notes": "Reports lag terms, cumulative lag effects, significance, N, and adjusted R2 to support the weak and unstable CPI transmission conclusion.",
        },
    ]
    return pd.DataFrame(rows)


def write_captions_file():
    lines = [
        "S1 Table. Variable names and abbreviations.",
        "S2 Table. Descriptive statistics of main variables for the long sample.",
        "S3 Table. Descriptive statistics of main variables for the short sample.",
        "S4 Table. Unit root test results for level series.",
        "S5 Table. Unit root test results for first-differenced series.",
        "S6 Table. Cumulative lag effects of the RMB-denominated oil price.",
        "S7 Table. Impacts of the RMB-denominated oil price on the crude oil import unit price.",
        "S8 Table. Impacts of import costs and upstream/midstream prices on CPI YoY.",
        "",
        "Notes for S4-S5: core_cpi_proxy_yoy is treated as a supplementary variable and is excluded from the core unit-root-test sheets because it is not a core ARDL variable used for the manuscript's main premise checks.",
        "Notes for S6: cumulative effects are estimated from the manuscript Table 4 specification, namely a 6-period DLM with oil_rmb lags 0-6 and indust_va_cum_yoy as the industrial value-added control.",
        "Notes for S8: the table reports lag-specific and cumulative CPI regressions for import costs and upstream/midstream price indicators; the pattern of small, mixed, and mostly unstable coefficients supports the manuscript conclusion that CPI transmission is weak, indirect, lagged, and unstable.",
    ]
    CAPTIONS_FILE.write_text("\n".join(lines) + "\n", encoding="utf-8")


def cpi_cost_pressure_table(df):
    df = add_industrial_control(df)
    yvar = "cpi_yoy"
    control_var = "indust_va_cum_yoy"
    model_specs = [
        ("Import cost to CPI", ["import_crude_oil_unit_price_yuan_per_kg"]),
        ("Fuel and power purchase price to CPI", ["ppi_fuel_power_yoy"]),
        ("PPI and fuel/power jointly to CPI", ["ppi_yoy", "ppi_fuel_power_yoy"]),
    ]
    rows = []
    for model_name, xvars in model_specs:
        model_result = fit_lagged_model(df, yvar, xvars, lags=6, controls=[control_var], min_obs=100)
        if model_result is None:
            continue
        model = model_result["fit"]
        predictor_cols = model_result["predictor_cols"]
        for xvar in xvars:
            for lag in range(7):
                term = f"{xvar}_lag{lag}"
                idx = 1 + predictor_cols.index(term)
                coef = float(model["beta"][idx])
                se = float(model["se"][idx])
                p_value = float(model["pvalues"][idx])
                rows.append(
                    {
                        "Model": model_name,
                        "Dependent variable": yvar,
                        "Explanatory variable": xvar,
                        "Term": f"{xvar} lag {lag}",
                        "Coefficient": round(coef, 4),
                        "Std. error": round(se, 4),
                        "95% CI lower": round(coef - 1.96 * se, 4),
                        "95% CI upper": round(coef + 1.96 * se, 4),
                        "p-value": round(p_value, 4),
                        "Significance": stars(p_value),
                        "Control variable": control_var,
                        "Adjusted R2": round(float(model["adj_r2"]), 4),
                        "N": int(model["nobs"]),
                    }
                )
            for horizon in (3, 6):
                cumulative = cumulative_row(
                    model_result,
                    xvar,
                    horizon,
                    f"Cumulative {xvar} effect, lags 0-{horizon}",
                )
                cumulative.update(
                    {
                        "Model": model_name,
                        "Dependent variable": yvar,
                        "Explanatory variable": xvar,
                        "Control variable": control_var,
                    }
                )
                rows.append(cumulative)
    return pd.DataFrame(rows)


def format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 42)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    df_long = load_master(LONG_FILE)
    df_short = load_master(SHORT_FILE)

    tables = {
        "S1 Table": variable_table(),
        "S2 Table": descriptive_stats(df_long, LONG_MAIN_VARS),
        "S3 Table": descriptive_stats(df_short, SHORT_MAIN_VARS),
        "S4 Table": adf_table(df_long, UNIT_ROOT_VARS, difference=False),
        "S5 Table": adf_table(df_long, UNIT_ROOT_VARS, difference=True),
        "S6 Table": cumulative_effect_table(df_short),
        "S7 Table": oil_to_import_price_table(df_short),
        "S8 Table": cpi_cost_pressure_table(df_short),
    }

    with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as writer:
        for sheet_name, table in tables.items():
            table.to_excel(writer, sheet_name=sheet_name, index=False)

    format_workbook(OUT_FILE)
    supporting_table_mapping().to_csv(MAPPING_FILE, index=False, encoding="utf-8-sig")
    write_captions_file()
    print(f"Supporting information workbook written to: {OUT_FILE}")
    print(f"Supporting table mapping written to: {MAPPING_FILE}")
    print(f"Suggested captions written to: {CAPTIONS_FILE}")


if __name__ == "__main__":
    main()
