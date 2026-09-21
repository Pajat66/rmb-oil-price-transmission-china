
import pandas as pd
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW = PROJECT_ROOT / "01_raw"

MASTER_CANDIDATES = [
    PROJECT_ROOT / "master_monthly_v1.xlsx",
    PROJECT_ROOT / "02_clean" / "monthly_master" / "master_monthly_v1.xlsx",
    PROJECT_ROOT / "02_clean" / "master_monthly_v1.xlsx",
    RAW / "master_monthly_v1.xlsx",
]

def first_existing(paths):
    for p in paths:
        if p.exists():
            return p
    return paths[0]

MASTER_FILE = first_existing(MASTER_CANDIDATES)

FILES = {
    "master": MASTER_FILE,
    "eia": RAW / "01_eia_brent" / "EIA_brent_usd_monthly_1987-latest_2026.3.22.csv",
    "fred": RAW / "02_fx_pbc_fred" / "FRED_exchus_cny_per_usd_monthly_1981-latest_2026.3.22.xlsx",
    "cpi_98_15": RAW / "03_nbs_price" / "98年-15年月度消费价格分类指数.xlsx",
    "cpi_16_20": RAW / "03_nbs_price" / "16年-20年月度消费价格分类指数.xlsx",
    "cpi_21_25": RAW / "03_nbs_price" / "21年-25年月度消费价格分类指数.xlsx",
    "cpi_26": RAW / "03_nbs_price" / "26年1-2月月度消费价格分类指数.xlsx",
    "ppi": RAW / "03_nbs_price" / "98年-26年2月工业生产者出厂价格指数.csv",
    "ppi_purchase": RAW / "03_nbs_price" / "98年-26年2月工业生产者购入价格指数.csv",
    "customs": RAW / "04_customs_trade" / "15年-26年2月月度原油进口量值数据.xlsx",
}


def ensure_files_exist():
    missing = [str(FILES[k]) for k in FILES if not Path(FILES[k]).exists()]
    if missing:
        msg = "\n".join(missing)
        raise FileNotFoundError("以下文件不存在，请检查路径：\n" + msg)


def ym_to_str(s):
    if pd.isna(s):
        return None
    s = str(s).strip()
    m = re.match(r"(\d{4})年(\d{1,2})月", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    try:
        return pd.to_datetime(s).strftime("%Y-%m")
    except Exception:
        return None


def read_wide_excel_xlsx(file_path: Path):
    df = pd.read_excel(file_path, header=None, engine="openpyxl")
    headers = df.iloc[2].tolist()
    df2 = df.iloc[3:].copy()
    df2.columns = headers
    return df2[df2["指标"].notna()].copy()


def read_wide_csv(file_path: Path):
    df = pd.read_csv(file_path, encoding="gbk", skiprows=2)
    return df[df["指标"].notna()].copy()


def extract_row(df, labels):
    if isinstance(labels, str):
        labels = [labels]
    for label in labels:
        row = df[df["指标"] == label]
        if not row.empty:
            row = row.iloc[0]
            out = []
            for col in df.columns[1:]:
                month = ym_to_str(col)
                if month is None:
                    continue
                val = row[col]
                if pd.notna(val):
                    out.append((month, val))
            return pd.DataFrame(out, columns=["month", "value"]).sort_values("month")
    return pd.DataFrame(columns=["month", "value"])


def build_master():
    months = pd.period_range("2015-01", "2026-02", freq="M").strftime("%Y-%m")
    master = pd.DataFrame({"month": months})

    # 1) Brent
    eia = pd.read_csv(FILES["eia"], skiprows=4)
    eia["month"] = pd.to_datetime(eia["Month"], format="%b %Y").dt.strftime("%Y-%m")
    price_col = [c for c in eia.columns if c != "Month"][0]
    eia = eia[["month", price_col]].rename(columns={price_col: "brent_usd"})
    master = master.merge(eia, on="month", how="left")

    # 2) Exchange rate
    fred = pd.read_excel(FILES["fred"], sheet_name="Monthly", engine="openpyxl")
    fred["month"] = pd.to_datetime(fred["observation_date"]).dt.strftime("%Y-%m")
    fred = fred[["month", "EXCHUS"]].rename(columns={"EXCHUS": "cny_per_usd"})
    master = master.merge(fred, on="month", how="left")

    # 3) CPI
    cpi_dfs = {
        "cpi_98_15": read_wide_excel_xlsx(FILES["cpi_98_15"]),
        "cpi_16_20": read_wide_excel_xlsx(FILES["cpi_16_20"]),
        "cpi_21_25": read_wide_excel_xlsx(FILES["cpi_21_25"]),
        "cpi_26": read_wide_excel_xlsx(FILES["cpi_26"]),
    }

    labelsets = {
        "cpi_yoy": [
            ["居民消费价格指数(上年同月=100)"],
        ],
        "cpi_transport_comm_yoy": [
            ["交通和通讯类居民消费价格指数(上年同月=100)"],
            ["交通和通信类居民消费价格指数(上年同月=100)"],
            ["交通通信类居民消费价格指数(上年同月=100)"],
        ],
        "core_cpi_proxy_yoy": [
            ["不包括食品和能源居民消费价格指数(上年同月=100)"],
        ],
    }

    for out_col, alias_groups in labelsets.items():
        parts = []
        for _, df in cpi_dfs.items():
            found = pd.DataFrame(columns=["month", "value"])
            for labels in alias_groups:
                found = extract_row(df, labels)
                if not found.empty:
                    break
            if not found.empty:
                parts.append(found)
        if parts:
            series = pd.concat(parts, ignore_index=True).drop_duplicates("month", keep="first").sort_values("month")
            series = series.rename(columns={"value": out_col})
            master = master.merge(series, on="month", how="left")
        else:
            master[out_col] = pd.NA

    # 4) PPI
    ppi_df = read_wide_csv(FILES["ppi"])
    ppi_purchase_df = read_wide_csv(FILES["ppi_purchase"])

    ppi = extract_row(ppi_df, "工业生产者出厂价格指数(上年同月=100)").rename(columns={"value": "ppi_yoy"})
    ppi_purchase = extract_row(ppi_purchase_df, "工业生产者购进价格指数(上年同月=100)").rename(columns={"value": "ppi_purchase_yoy"})
    ppi_fuel = extract_row(ppi_purchase_df, "燃料、动力类购进价格指数(上年同月=100)").rename(columns={"value": "ppi_fuel_power_yoy"})

    master = master.merge(ppi, on="month", how="left")
    master = master.merge(ppi_purchase, on="month", how="left")
    master = master.merge(ppi_fuel, on="month", how="left")

    # 5) Customs crude-oil import data
    customs = pd.read_excel(FILES["customs"], engine="openpyxl")
    customs["商品编码"] = pd.to_numeric(customs["商品编码"], errors="coerce")
    crude = customs[customs["商品编码"] == 27090000].copy()
    crude["month"] = pd.to_datetime(crude["数据年月"].astype(str), format="%Y%m").dt.strftime("%Y-%m")
    crude["import_crude_oil_qty_10k_ton"] = pd.to_numeric(crude["第一数量"], errors="coerce") / 1e7
    crude["import_crude_oil_value_100m_rmb"] = pd.to_numeric(crude["人民币"], errors="coerce") / 1e8
    crude["import_crude_oil_unit_price_yuan_per_kg"] = (
        pd.to_numeric(crude["人民币"], errors="coerce") / pd.to_numeric(crude["第一数量"], errors="coerce")
    )
    crude = crude[[
        "month",
        "import_crude_oil_qty_10k_ton",
        "import_crude_oil_value_100m_rmb",
        "import_crude_oil_unit_price_yuan_per_kg",
    ]]
    master = master.merge(crude, on="month", how="left")

    # 6) Derived variables
    master["oil_rmb"] = master["brent_usd"] * master["cny_per_usd"]

    ordered_cols = [
        "month",
        "brent_usd",
        "cny_per_usd",
        "oil_rmb",
        "cpi_yoy",
        "cpi_transport_comm_yoy",
        "core_cpi_proxy_yoy",
        "ppi_yoy",
        "ppi_purchase_yoy",
        "ppi_fuel_power_yoy",
        "import_crude_oil_qty_10k_ton",
        "import_crude_oil_value_100m_rmb",
        "import_crude_oil_unit_price_yuan_per_kg",
    ]
    return master[ordered_cols]


def write_to_excel(df, output_file: Path):
    wb = load_workbook(FILES["master"])
    ws = wb[wb.sheetnames[0]]

    # Ensure the header exists and matches the dataframe columns
    expected_headers = list(df.columns)
    for c, h in enumerate(expected_headers, start=1):
        ws.cell(1, c, value=h)

    # Clear the previous data area
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=max(ws.max_column, len(expected_headers))):
        for cell in row:
            cell.value = None

    # Write the updated data
    for r, (_, row) in enumerate(df.iterrows(), start=2):
        for c, value in enumerate(row.tolist(), start=1):
            ws.cell(r, c, value=value)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(expected_headers) + 1):
        letter = get_column_letter(col_idx)
        values = [ws.cell(row=r, column=col_idx).value for r in range(1, min(ws.max_row, 20) + 1)]
        max_len = max(len(str(v)) if v is not None else 0 for v in values)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 28)

    wb.save(output_file)


if __name__ == "__main__":
    ensure_files_exist()
    master = build_master()
    output = FILES["master"].with_name("master_monthly_v1_filled_full_openpyxl.xlsx")
    write_to_excel(master, output)
    print("项目根目录：", PROJECT_ROOT)
    print("主表模板：", FILES["master"])
    print("已生成：", output)
