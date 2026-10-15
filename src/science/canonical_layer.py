"""Historical source representations and oil builder, reconstructed from raw data.
Default openpyxl serialization is retained; no rounding or precision overrides.
"""
from pathlib import Path
from openpyxl import Workbook
import numpy as np
import pandas as pd

def validate_required_columns(df, required, label='dataset'):
    absent = sorted(set(required) - set(df.columns))
    if absent:
        raise ValueError(f'{label}: missing required columns: {absent}')
    return True

def build_oil_rmb(df, label='dataset', tolerance=1e-8):
    """Preserve existing values; derive only blank values and report discrepancies."""
    validate_required_columns(df, ['brent_usd', 'cny_per_usd'], label)
    out = df.copy()
    expected = out.brent_usd * out.cny_per_usd
    source = out.oil_rmb.copy() if 'oil_rmb' in out else pd.Series(np.nan, index=out.index)
    both = source.notna() & expected.notna()
    diffs = (source[both] - expected[both]).abs()
    mismatch = int((diffs > tolerance).sum())
    report = {
        'dataset': label, 'source_nonmissing': int(source.notna().sum()),
        'derived_count': int((source.isna() & expected.notna()).sum()),
        'source': 'derived variable' if source.notna().sum() == 0 else 'preserved existing; blanks derived',
        'max_absolute_difference': float(diffs.max()) if len(diffs) else None,
        'mean_absolute_difference': float(diffs.mean()) if len(diffs) else None,
        'mismatch_count': mismatch,
        'comparison_tolerance': tolerance,
    }
    out['oil_rmb_source'] = np.where(source.notna(), 'source', 'derived')
    out['oil_rmb'] = source.where(source.notna(), expected)
    if mismatch:
        raise ValueError(f'{label}: {mismatch} source oil_rmb values disagree with Brent × FX: {report}')
    return out, report

def write_master_source(frame,path,formula_oil=False):
    """Recreate historical cell values/formulas using the original writer's defaults.

    This is a source-format artifact; the numerical long analysis master is kept.
    The original short writer writes numeric oil prices. The long writer writes
    IF/OR formula cells without caches. Both are read with data_only semantics.
    """
    wb=Workbook();ws=wb.active
    ws.title='master_monthly_long' if formula_oil else 'master_monthly_short'
    ws.append(frame.columns.tolist())
    oil=frame.columns.get_loc('oil_rmb')+1
    assert oil==4 and frame.columns[:3].tolist()==['month','brent_usd','cny_per_usd']
    for i,values in enumerate(frame.itertuples(index=False,name=None),start=2):
        for j,value in enumerate(values,start=1):
            if formula_oil and j==oil:
                ws.cell(i,j,value=f'=IF(OR(B{i}="",C{i}=""),"",B{i}*C{i})')
            else:ws.cell(i,j,value=None if pd.isna(value) else value)
    wb.save(path);wb.close()
    return path

def read_historical_stage(source_path,clean_path):
    """Mirror chapter2 clean_master and its XLSX write/read boundary.

    Order and calendar are validated by the caller. No values are filled.
    No frozen workbook is opened or used as a template.
    """
    frame=pd.read_excel(source_path,engine='openpyxl')
    frame['month']=pd.to_datetime(frame['month'],errors='coerce')
    frame=frame.sort_values('month').reset_index(drop=True)
    frame=frame.drop_duplicates(subset=['month'],keep='first')
    for col in frame.columns:
        if col!='month':frame[col]=pd.to_numeric(frame[col],errors='coerce')
    frame.to_excel(clean_path,index=False,engine='openpyxl')
    return pd.read_excel(clean_path,engine='openpyxl')
